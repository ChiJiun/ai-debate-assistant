from __future__ import annotations

from io import BytesIO
from typing import Mapping

from docx import Document


def _add_markdownish_text(document: Document, text: str) -> None:
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line[:3].replace(".", "").isdigit() and ". " in line[:5]:
            document.add_paragraph(line, style="List Number")
        else:
            document.add_paragraph(line)


def build_docx(settings: Mapping[str, str], sections: Mapping[str, str]) -> BytesIO:
    document = Document()
    document.add_heading("辯論助理準備資料", level=0)

    document.add_heading("基本設定", level=1)
    for label, value in settings.items():
        document.add_paragraph(f"{label}: {value}")

    for title, content in sections.items():
        document.add_heading(title, level=1)
        _add_markdownish_text(document, content)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
